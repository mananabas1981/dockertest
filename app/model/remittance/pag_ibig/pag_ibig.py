from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Border, Side, Font 
from os import sep, path
from io import BytesIO
from base64 import b64encode

class PagIbig01:
  def __init__(self, data = {}):

    self.data = data

    self.excel_template = f'app{sep}model{sep}remittance{sep}pag_ibig{sep}pag_ibig_01.xlsx'

    if not path.exists(self.excel_template):
      print(self.excel_template + ' not found!')
      raise ValueError(self.excel_template)


  def render(self):


    workbook = load_workbook(filename = self.excel_template)

    # # Select the active worksheet
    ws = workbook.worksheets[0] # worksheet number 

    # Data can be assigned directly to cells
    # ws['A1'] = 42

    # # Rows can also be appended
    # ws.append([1, 2, 3])

    # # Python types will automatically be converted
    # import datetime
    # ws['A2'] = datetime.datetime.now()

    workbook.save(filename = f'tmp{sep}test.xlsx')
    workbook.close()

  def output(self):
    pass