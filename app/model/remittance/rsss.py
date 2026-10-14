from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Border, Side, Font 

from os import sep
from io import BytesIO
from base64 import b64encode

class SSS():
  def __init__(self, ids, template = ''):
    self.template = template
    self.data = {}
    self.item = []
    self.code = ''
    self.ids = ''
    self.collect = []

  def set_id(self, ids):
    self.ids = ids

  def set_code(self, code):
    self.code = code

  def set_data(self, _data):
    self.data = _data 


  def set_item(self, _item):
    self.item = _item

  def set_collect(self, collect):
    self.collect = collect




  def contribution(self, ws, collect):
    # FOR CONTRIBUTIONS

    border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )
    cnt = 7

    for item in collect:
      ws[f'C{cnt}'] = item['card_id']
      ws[f'D{cnt}'] = item['lname'].upper()
      ws[f'E{cnt}'] = item['fname'].upper()
      ws[f'F{cnt}'] = item['ename'].upper()
      ws[f'G{cnt}'] = item['mname'].upper()
      ws[f'H{cnt}'] = item['dob'].upper()


      ws[f'C{cnt}'].border = border
      ws[f'D{cnt}'].border = border
      ws[f'E{cnt}'].border = border
      ws[f'F{cnt}'].border = border
      ws[f'G{cnt}'].border = border
      ws[f'H{cnt}'].border = border


      ws[f'I{cnt}'] = item['percov']      

      ws[f'I{cnt}'].border = border
      ws[f'K{cnt}'].border = border
      ws[f'L{cnt}'].border = border
      
      
      ee_share = item['value']
      range_t = item['range']

      if len(ee_share) == 1:
        for item in ee_share[0]:
          name, value, type_t = item

          if type_t == 'c':
            if name.upper() == 'CONTRIBUTION':
              ws[f'M{cnt}'] = value
              ws[f'M{cnt}'].border = border

            if name.upper() == 'EC':
              ws[f'J{cnt}'] = value
              ws[f'J{cnt}'].border = border            


      elif len(ee_share) == 2:
        # m = ord('M')
        gvalue = 0

        for j in range(len(ee_share)):
          e_share = ee_share[j]
          for k in range(len(e_share)):
            item = e_share[k]
            name, value, type_t = item
            if type_t == 'c':
              gvalue += value
              if name.upper() == 'CONTRIBUTION':
                ws[f'M{cnt}'] = gvalue
                ws[f'M{cnt}'].border = border
                ws[f'{chr(m)}{cnt}'] = f'{value} <-> ({range_t[j][0]} - {range_t[j][1]})'
                m += 1

      cnt += 1

      ws.protection.sheet = True
      ws.protection.password = self.ids
      ws.protection.enable()
      # END FOR CONTRIBUTIONS

  # def loan(self, ws, collect):
  #   cnt = 5
  #   for item in collect:
  #     ws[f'A{cnt}'] = item['card_id']
  #     ws[f'C{cnt}'] = item['lname'].upper()
  #     ws[f'D{cnt}'] = item['fname'].upper()
  #     ws[f'E{cnt}'] = item['ename'].upper()
  #     ws[f'F{cnt}'] = item['mname'].upper()

       
  #     ee_share = item['value']
  #     range_t = item['range']

  #     if len(ee_share) == 1:
  #       for item in ee_share[0]:
  #         name, value, type_t = item

  #         if type_t == 'l':
  #           if name.upper() == 'MPL':
  #             ws[f'G{cnt}'] = name.upper() # loan type
  #             ws[f'H{cnt}'] = value
              


  #     elif len(ee_share) == 2:
  #       m = ord('J')
  #       gvalue = 0

  #       for j in range(len(ee_share)):
  #         e_share = ee_share[j]
  #         for k in range(len(e_share)):
  #           item = e_share[k]
  #           name, value, type_t = item
  #           if type_t == 'l':
  #             gvalue += value
  #             if name.upper() == 'MPL':
  #               ws[f'G{cnt}'] = name.upper()  # loan type

  #               ws[f'H{cnt}'] = gvalue
  #               ws[f'{chr(m)}{cnt}'] = f'{value} <-> ({range_t[j][0]} - {range_t[j][1]})'
  #               m += 1

  #     cnt += 1

  #     ws.protection.sheet = True
  #     ws.protection.password = self.ids
  #     ws.protection.enable()


  def generate(self):


    workbook = load_workbook(filename = self.template)

    # # Select the active worksheet
    
    # Data can be assigned directly to cells
    # ws['A1'] = 42

    # # Rows can also be appended
    # ws.append([1, 2, 3])

    # # Python types will automatically be converted
    # import datetime
    # ws['A2'] = datetime.datetime.now()
    ws = workbook.worksheets[0] # worksheet number 
    self.contribution(ws, self.collect)
    # ws = workbook.worksheets[1] # second shets
    # self.loan(ws, self.collect)

    bytes_t = BytesIO()

    workbook.save(bytes_t)
    workbook.close()

    ret = b64encode(bytes_t.getvalue()).decode()

    bytes_t.close()

    return ret
    

  def download(self):
    return { 
      'name': f'{self.ids}.xlsx',
      'data': self.generate()
    }
