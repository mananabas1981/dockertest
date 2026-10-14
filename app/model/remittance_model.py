from app.store.payroll_store import PayrollStore
from app.util.json import json_parse
from app.store.client_store import ClientStore
from app.store.deduction_store import DeductionStore
from app.store.default_ded import DefaultDeduction
from app.store.card_store import CardStore

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
from os import sep

from io import BytesIO
from base64 import b64encode
from app.util.date import date_time


class RemittanceModel:
    def __init__(this, con):
        this._con = con

    def generate(this, refid, key, all_include_zero = False):
        
        pstore = PayrollStore(this._con)
        
        items = pstore.get_item(refid)
        
        this._con.close()

        # early close to be able to re-use the connection from other connection
        # we use this method to avoid server overloaded!

        lists = []

        for item in items:
            deds = json_parse(item['deduction'])
            emp_info = json_parse(item['emp_info'])

            info = {
                "name": emp_info,
                "item": []
            }
            
            has_value = False
            for ded in deds:
                if (ded['code'] == key):
                    info["item"].append(ded)
                    if float(ded["value"]) > 0.0:
                        has_value = True

            if not all_include_zero:
                if has_value:
                    lists.append(info)
            else:
                lists.append(info)


        # print(rem.refid)
        # print(rem.key)
        return lists


# Upload Model
class RemittanceUploadModel:
    def __init__(self, con = None, job_status = 'j'):
        self.job_status = job_status



# Download Model

class RemittanceDownloadModel:
    def __init__(self, con = None, job_status = 'j'):

        fn = {
            'j': 'JobOrder',
            'c': 'ContractOfService',
            'r': 'Regular'
        }

        self.msg = {
            'ok': False, 
            'msg': '',
            'data': '',
            'file_name': f'{fn[job_status]}-{date_time()}'
        }


        self.con = con
        cs = ClientStore(self.con)
        clients = cs._getClients(emp_status = job_status, limit = None, is_active = 'y')

        dfd_obj = DefaultDeduction(self.con) 
        card_obj = CardStore(self.con)

        self.persons = []
        for item in clients:
            
            pos_info = item['pos_info']

            rate = ''
            if pos_info: 
                for _pos_item in pos_info:
                    if _pos_item['key'] == 'Rate':
                        rate = _pos_item['value']
                        break
            
            # default deduction 
            dfd = dfd_obj.get_default_by_emp_id(item['id'])

            # end default deduction

            # card information like pag-ibig, sss no.
            card_data = card_obj.get_by_emp_id(item['id'])

            # end card

            self.persons.append({
                'id': item['id'],
                'hr_id': item['hr_id'],
                'name': f'''{item['lname']}, {item['fname']} {item['ename']} {item['mname']}'''.upper(),
                'rate': rate,
                'defaul_ded_info': dfd,
                'card': card_data 
            })


        # get all the deduction base on job_status
        ded_store = DeductionStore(self.con)

        ded_data = ded_store.getDeduction(status = job_status) 

        self.deds = []


        for item in ded_data:
            code, name = item['code'], item['name']
            self.deds.append({
                'code': code,
                'name': name
            })

        # print(self.persons)
        self.create_html()
    
    def json_format(self):
        ret = {
            "base": {}
        }
        
        base = {}
        for item in self.deds:
            if not item["code"] in base:
                base[item["code"]] = []
            
            base[item["code"]].append(item["name"])

        
        ret["base"] = base


        return ret



    def create_html(self):

        

        res = self.json_format()
        base = res["base"]

        # create a button for the base
        styles = []
        styles.append("""
            .btn {
                color: white;
                background-color: #5C6BC0;
                cursor: pointer;
                border: 1px solid rgba(0, 0, 0, 0.3);
                padding-left: 2px;
                padding-right: 2px;
                display: inline-block;
                
            }
        """)

        print(base)

        buttons = []
        for item in base:
            buttons.append(f"""
                <button class="btn">
                    {item}
                </button>
            """)
            
        template = []
        # create base htmls




        script = """


            const buttons = document.getElementsByClassName("btn");
            const btn_len = buttons.length;
            for (let i = 0; i < btn_len; i++) {
                buttons[i].addEventListener("click", () => {
                    console.log(i);
                });
            }
            
        """

        # self.msg["json"] = res
        self.msg['data'] = "<style>" + "".join(styles) \
            + "</style>" + "".join(buttons) \
            + f"<script>{script}</script>"

        

    def create_excel(self):
        wb = Workbook()

        # ws = wb.create_sheet('test1', 0)
        # ws = wb.create_sheet('test2', 1)
        center_style = Alignment(horizontal='center', vertical='center')
        border_style = Side(style='thin', color='000000')
        red_font = Font(color='FF0000', bold = True)

        bordered = Border(left = border_style, right = border_style, top = border_style, bottom = border_style)
        
        bold_font = Font(bold = True)
        fill = PatternFill(start_color='FF0000', fill_type='solid')
        index = 0


        base = {}
        for item in self.deds:
            if not item["code"] in base:
                base[item["code"]] = []
            
            base[item["code"]].append(item["name"])



        ################################################################################
        #   DRAWING HEADERS
        ################################################################################
        for item in base:

            code = item

            s = f"{code}"
            s = "".join(c for c in s if c.isalnum())

            ws = wb.create_sheet(s, index)

            # processing here per sheets

            ws['A1'] = 'CODE'
            ws['A2'] = 'TYPE' 

            ws['B1'] = code
            ws['B2'] = "---"

            # headers
            ws['A4'] = 'SYSTEM ID'
            ws['B4'] = 'HRMS ID'
            ws['C4'] = 'NAME'
            ws['D4'] = "SALARY"
            ws['E4'] = f'{code} No.'
            

            # ws['F4'] = '1ST HALF'
            # ws['G4'] = '2ND HALF'
            # ws['H4'] = 'WHOLE MONTH'



            # generate the all header here

            
            

            #ws.merge_cells("F3:H3")
            # base headers
            j = ["", ord("F")]
            ded_names = base[code]
            for ded_name in ded_names:
                
                merge = f"{j[0]}{chr(j[1])}3:{j[0]}{chr(j[1] + 2)}3"

                ws.merge_cells(merge)

                ws[f"{j[0]}{chr(j[1])}3"] = ded_name
                ws[f"{j[0]}{chr(j[1])}3"].font = red_font
                ws[f"{j[0]}{chr(j[1])}3"].alignment = center_style


                # here the value 
                
                ws.column_dimensions[f"{j[0]}{chr(j[1])}"].width     = 11
                



                # starting here changing the value of the coordinate
                
                extra = []

                if (j[1] + 3) >= 90:
        
 
                    k = j[1]
                    while k <= j[1] + 3:
                        if k <= 90:
                            extra.append([j[0], chr(k)])
                        k += 1
                        
                    j[1] = ord("A")

                    if j[0]:
                        j[0] = chr(ord("A") + 1)
                    else:
                        j[0] = 'A'

                # print(extra)

                zero = extra[1] if len(extra) > 0 else [j[0], chr(j[1])]
                one  = extra[1] if len(extra) > 1 else [j[0], chr(j[1] + 1)]
                two  = extra[2] if len(extra) > 2 else [j[0], chr(j[1] + 2)]
                

                ws[f"{zero[0]}{zero[1]}4"] = "1ST HALF"
                ws[f"{one[0]}{one[1]}4"] = "2ND HALF"
                ws[f"{one[0]}{two[1]}4"] = "WHOLE"



                ws[f"{two[0]}{two[1]}4"].border = Border(right = Side(style = 'thick', color = '000000'))
                
                ws.column_dimensions[f"{one[0]}{one[1]}"].width = 11
                ws.column_dimensions[f"{two[0]}{two[1]}"].width = 11
                
                ws.column_dimensions[f"{two[0]}{two[1]}"].border = Border(right = Side(style = 'thick', color = '000000'))
                # ws.column_dimensions[f"{chr(j)}"].fill     = fill
                # ws.column_dimensions[f"{chr(j + 1)}"].fill = fill
                # ws.column_dimensions[f"{chr(j + 2)}"].fill = fill
                
    

                j[1] += 3

            # value of the headers:
            
        ################################################################################
        #   END DRAWING HEADERS
        ################################################################################        


            # ws['A4'].font = bold_font
            # ws['B4'].font = bold_font
            # ws['C4'].font = bold_font
            # ws['D4'].font = bold_font
            # ws['E4'].font = bold_font
            # ws['F4'].font = bold_font
            # ws['G4'].font = bold_font
            # ws['H4'].font = bold_font
            # ws['I4'].font = bold_font

            

            
            # ws['F4'].fill = fill
            # ws['G4'].fill = fill
            # ws['H4'].fill = fill

            # ws['D4'].fill = fill
            # ws['D4'].font = Font(bold = True, color = 'FFFFFF')

            # ws['I4'].fill = fill
            # ws['I4'].font = Font(bold = True, color = 'FFFFFF')

            # ws['F4'].font = Font(bold = True, color = 'FFFFFF')
            # ws['G4'].font = Font(bold = True, color = 'FFFFFF')
            # ws['H4'].font = Font(bold = True, color = 'FFFFFF')

            # body

            ws.freeze_panes = 'F5'
            

            ws.column_dimensions['A'].width = 10
            ws.column_dimensions['C'].width = 10
            ws.column_dimensions['C'].width = 20
            ws.column_dimensions['D'].width = 18
            ws.column_dimensions['E'].width = 15


            i = 5
            p = 1
            for person in self.persons:
                ws[f'C{i}'].font = bold_font

                ws[f'A{i}'] = person['id']
                ws[f'B{i}'] = person['hr_id']
                ws[f'C{i}'] = person['name']
                
                ws[f'D{i}'] = person['rate']


                print(p, person["name"])
                p += 1
                #############################################################
                #   DRAW ALL THE VALUE OF THE INPUT
                #############################################################
                for item in base:

                    code = item

                    

                    s = f"{code}"
                    s = "".join(c for c in s if c.isalnum())

                    ws = wb.create_sheet(s, index)

                    # processing here per sheets

                    ws[f'E{i}'] = f'00000'
                    

                    # ws['F4'] = '1ST HALF'
                    # ws['G4'] = '2ND HALF'
                    # ws['H4'] = 'WHOLE MONTH'



                    # generate the all header here

                    
                    

                    #ws.merge_cells("F3:H3")
                    # base headers
                    j = ["", ord("F")]
                    ded_names = base[code]
                    for ded_name in ded_names:
                        
                        merge = f"{j[0]}{chr(j[1])}3:{j[0]}{chr(j[1] + 2)}3"

                        ws.merge_cells(merge)

                        ws[f"{j[0]}{chr(j[1])}3"] = ded_name
                        ws[f"{j[0]}{chr(j[1])}3"].font = red_font
                        ws[f"{j[0]}{chr(j[1])}3"].alignment = center_style


                        # here the value 
                        
                        ws.column_dimensions[f"{j[0]}{chr(j[1])}"].width     = 11
                        



                        # starting here changing the value of the coordinate
                        
                        extra = []

                        if (j[1] + 3) >= 90:
                
        
                            k = j[1]
                            while k <= j[1] + 3:
                                if k <= 90:
                                    extra.append([j[0], chr(k)])
                                k += 1
                                
                            j[1] = ord("A")

                            if j[0]:
                                j[0] = chr(ord("A") + 1)
                            else:
                                j[0] = 'A'

                        # print(extra)

                        zero = extra[1] if len(extra) > 0 else [j[0], chr(j[1])]
                        one  = extra[1] if len(extra) > 1 else [j[0], chr(j[1] + 1)]
                        two  = extra[2] if len(extra) > 2 else [j[0], chr(j[1] + 2)]
                        

                        ws[f"{zero[0]}{zero[1]}4"] = "1ST HALF"
                        ws[f"{one[0]}{one[1]}4"] = "2ND HALF"
                        ws[f"{one[0]}{two[1]}4"] = "WHOLE"



                        ws[f"{two[0]}{two[1]}4"].border = Border(right = Side(style = 'thick', color = '000000'))
                        
                        ws.column_dimensions[f"{one[0]}{one[1]}"].width = 11
                        ws.column_dimensions[f"{two[0]}{two[1]}"].width = 11
                        
                        ws.column_dimensions[f"{two[0]}{two[1]}"].border = Border(right = Side(style = 'thick', color = '000000'))
                        # ws.column_dimensions[f"{chr(j)}"].fill     = fill
                        # ws.column_dimensions[f"{chr(j + 1)}"].fill = fill
                        # ws.column_dimensions[f"{chr(j + 2)}"].fill = fill
                        
            

                        j[1] += 3

                    # value of the headers:

                #############################################################
                #   END DRAW ALL THE VALUE OF THE INPUT
                #############################################################


                # plotting all the value of the information

                # # initialized the default deduction
                # df_data = {
                #     '0': '',
                #     '1': '',
                #     '2': ''
                # }
                # term = ''

                # dd_info = person['defaul_ded_info']
                # for d_item in dd_info:
                #     if d_item['code'] == item['code'] and d_item['name'] == item['name']:
                #         if not df_data['0'] and d_item['type_ded'] == '0':
                #             df_data['0'] = d_item['value']
                        
                #         if not df_data['1'] and d_item['type_ded'] == '1':
                #             df_data['1'] = d_item['value']

                #         if not df_data['2'] and d_item['type_ded'] == '2':
                #             df_data['2'] = d_item['value']
                #         if not term:
                #             term = d_item['end_term']

                # ws[f'F{i}'] = df_data['1']
                # ws[f'G{i}'] = df_data['2']
                # ws[f'H{i}'] = df_data['0']

                # ws[f'F{i}'].font = Font(bold = True)
                # ws[f'G{i}'].font = Font(bold = True)
                # ws[f'H{i}'].font = Font(bold = True)

                # ws[f'I{i}'] = term
                # # end initialized of the default deduction


                # # initialized the card like pag-ibig number
                # c_number = ''
                # c_info = person['card']
                # for c_item in c_info:
                #     if c_item['code'] == item['code']: # matching the code card and the item deduction code
                #         c_number = c_item['card_no']

                # ws[f'D{i}'] = c_number
                # # end initialized card like pag-ibig number, sss number

                i += 1                

            # end processing per sheets
            index += 1

      
        byte_io = BytesIO()
        wb.save(byte_io)
        wb.close()

        self.msg['ok'] = True
        self.msg['data'] = b64encode(byte_io.getvalue()).decode()


    def output(self) -> dict:
        return self.msg


